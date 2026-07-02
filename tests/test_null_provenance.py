from types import SimpleNamespace

import openpyxl

from app.core.excel_mapping import CELL_PAGE, CELL_SOURCE_TYPES, EXCEL_MAPPING
from app.models import (
    DocumentExtraction,
    ExtractedTable,
    MeridianExtractionResponse,
    MeridianStructuredPage,
    PageMetadata,
    TableCell,
    TableRow,
)
from app.services.excel_populator import FORMULA_FILL, _apply_static_formula_cache_fixes
from app.services.excel_populator import _build_formula_record, _build_null_record
from app.services.excel_populator import _build_provenance_record, _is_missing_excel_value
from app.services.excel_populator import _resolve_formula_cells


def test_unclear_logic_cell_with_value_is_reported_as_inferred():
    # "unclear_logic" is a NullCategory, not a SourceType. When such a cell
    # actually carries a value it must be emitted as a valid SourceType
    # ("inferred") and never claim a PDF source — otherwise FieldProvenance
    # validation fails (regression: export-excel 422).
    doc_extraction = DocumentExtraction(
        document_id="doc",
        filename="sample.pdf",
        sha256="sha",
        page_count=1,
        pages=[_page_meta(1)],
        tables=[],
    )

    record = _build_provenance_record(
        "K26", 0.85, _structured_page_1(), doc_extraction, "sample.pdf"
    )

    assert record.source_type == "inferred"
    assert record.value == 0.85
    assert record.bbox is None


def _page_meta(page_number: int, classification: str = "vector_text") -> PageMetadata:
    return PageMetadata(
        page_number=page_number,
        width=100.0,
        height=100.0,
        rotation=0,
        classification=classification,
        text_blocks=1,
        text_lines=1,
        text_chars=1,
        image_count=0,
        drawing_count=0,
    )


def _table_with_blank_static_cell() -> ExtractedTable:
    return ExtractedTable(
        table_id="p1_t1",
        page_number=1,
        bbox=(0.0, 0.0, 100.0, 100.0),
        columns=[],
        rows=[
            TableRow(
                index=0,
                cells=[
                    TableCell(row=0, column=0, text=""),
                    TableCell(row=0, column=1, text=""),
                    TableCell(row=0, column=2, text=""),
                ],
            ),
            TableRow(
                index=1,
                cells=[
                    TableCell(row=1, column=0, text=""),
                    TableCell(row=1, column=1, text=""),
                    TableCell(row=1, column=2, text="", bbox=(10.0, 20.0, 30.0, 40.0)),
                ],
            ),
        ],
        confidence=1.0,
    )


def _table_with_cell(
    table_id: str,
    page_number: int,
    row_idx: int,
    col_idx: int,
    bbox: tuple[float, float, float, float],
) -> ExtractedTable:
    return ExtractedTable(
        table_id=table_id,
        page_number=page_number,
        bbox=(0.0, 0.0, 100.0, 100.0),
        columns=[],
        rows=[
            TableRow(
                index=row,
                cells=[
                    TableCell(
                        row=row,
                        column=col,
                        text="value" if row == row_idx and col == col_idx else "",
                        bbox=bbox if row == row_idx and col == col_idx else None,
                    )
                    for col in range(col_idx + 1)
                ],
            )
            for row in range(row_idx + 1)
        ],
        confidence=1.0,
    )


def _structured_page_1(header: dict | None = None) -> MeridianExtractionResponse:
    return MeridianExtractionResponse(
        document_id="doc",
        filename="sample.pdf",
        page_count=1,
        pages=[
            MeridianStructuredPage(
                page_number=1,
                page_type="gdc_estimation",
                header=header or {},
            )
        ],
    )


def test_cell_page_covers_every_non_default_mapping():
    missing = [
        coord
        for coord in EXCEL_MAPPING
        if CELL_SOURCE_TYPES.get(coord) not in ("default", "unclear_logic")
        and CELL_PAGE.get(coord) is None
    ]

    assert missing == []


def test_boolean_mapping_preserves_missing_value_as_null():
    extraction = SimpleNamespace(
        pages=[
            SimpleNamespace(),
            SimpleNamespace(),
            SimpleNamespace(resource_requirements={"setup_changeover_considered": None}),
            SimpleNamespace(),
            SimpleNamespace(
                page_type="assembly_estimation",
                assembly_resource_requirements={
                    "feasible_to_use_machining_operator_for_assembly": None
                },
                after_assembly_machining={
                    "involved": None,
                    "resource_requirements": {
                        "feasible_to_use_machining_cell_operator": None
                    },
                },
            ),
        ]
    )

    assert EXCEL_MAPPING["W29"](extraction) is None
    assert EXCEL_MAPPING["W33"](extraction) is None
    assert EXCEL_MAPPING["W52"](extraction) is None
    assert EXCEL_MAPPING["AI51"](extraction) is None


def test_blank_strings_are_treated_as_missing_excel_values():
    assert _is_missing_excel_value(None) is True
    assert _is_missing_excel_value("") is True
    assert _is_missing_excel_value("   ") is True
    assert _is_missing_excel_value(0) is False
    assert _is_missing_excel_value(False) is False


def test_null_record_uses_blank_source_cell_bbox_when_locatable():
    doc_extraction = DocumentExtraction(
        document_id="doc",
        filename="sample.pdf",
        sha256="sha",
        page_count=1,
        pages=[_page_meta(1)],
        tables=[_table_with_blank_static_cell()],
    )

    # Page 1 IS present — capital_investments proves real GDC content exists,
    # but rfq_no (C2) was left blank in the PDF.
    structured = MeridianExtractionResponse(
        document_id="doc",
        filename="sample.pdf",
        page_count=1,
        pages=[
            MeridianStructuredPage(
                page_number=1,
                page_type="gdc_estimation",
                header={"rfq_no": None, "customer": "CTT"},
                capital_investments=[{"description": "Core shooting M/c"}],
            )
        ],
    )
    record = _build_null_record("C2", structured, doc_extraction, "sample.pdf")

    assert record.source_type == "null"
    assert record.null_category == "data_absent"
    assert record.blocks_approval is False   # data_absent is informational, not blocking
    assert record.page_number == 1
    assert record.bbox == (10.0, 20.0, 30.0, 40.0)


def test_null_record_skips_static_source_for_missing_logical_page():
    doc_extraction = DocumentExtraction(
        document_id="doc",
        filename="sample.pdf",
        sha256="sha",
        page_count=3,
        pages=[_page_meta(3)],
        tables=[_table_with_cell("p3_t1", 3, 1, 1, (10.0, 20.0, 30.0, 40.0))],
    )
    structured = MeridianExtractionResponse(
        document_id="doc",
        filename="sample.pdf",
        page_count=3,
        pages=[
            MeridianStructuredPage(
                page_number=3,
                physical_page_number=None,
                page_type="machining_estimation",
                machining_operations=[],
            )
        ],
    )

    record = _build_null_record("AC10", structured, doc_extraction, "sample.pdf")

    assert record.source_type == "null"
    assert record.null_category == "page_missing"
    assert record.blocks_approval is True
    assert record.page_number is None
    assert record.bbox is None


def test_static_provenance_uses_shifted_physical_page():
    doc_extraction = DocumentExtraction(
        document_id="doc",
        filename="sample.pdf",
        sha256="sha",
        page_count=4,
        pages=[_page_meta(4)],
        tables=[_table_with_cell("p4_t1", 4, 8, 2, (11.0, 22.0, 33.0, 44.0))],
    )
    structured = MeridianExtractionResponse(
        document_id="doc",
        filename="sample.pdf",
        page_count=4,
        pages=[
            MeridianStructuredPage(
                page_number=5,
                physical_page_number=4,
                page_type="assembly_estimation",
            )
        ],
    )

    record = _build_provenance_record("W10", 3.5, structured, doc_extraction, "sample.pdf")

    assert record.source_type == "pdf_cell"
    assert record.page_number == 4
    assert record.bbox == (11.0, 22.0, 33.0, 44.0)


def test_null_record_marks_missing_page_from_page_metadata():
    doc_extraction = DocumentExtraction(
        document_id="doc",
        filename="sample.pdf",
        sha256="sha",
        page_count=1,
        pages=[_page_meta(1)],
        tables=[],
    )

    record = _build_null_record("AQ18", _structured_page_1(), doc_extraction, "sample.pdf")

    assert record.source_type == "null"
    assert record.null_category == "page_missing"
    assert record.blocks_approval is True


def test_null_record_marks_present_unparseable_page_as_extraction_failure():
    doc_extraction = DocumentExtraction(
        document_id="doc",
        filename="sample.pdf",
        sha256="sha",
        page_count=3,
        pages=[_page_meta(1), _page_meta(2), _page_meta(3, "image_only")],
        tables=[],
    )
    # Structured output includes a machining page with data so page is "present"
    # but the raw extraction shows it's image_only → extraction_failure.
    structured = MeridianExtractionResponse(
        document_id="doc",
        filename="sample.pdf",
        page_count=3,
        pages=[
            MeridianStructuredPage(
                page_number=1,
                page_type="gdc_estimation",
                header={"rfq_no": "123"},
            ),
            MeridianStructuredPage(
                page_number=3,
                page_type="machining_estimation",
                header={"rfq_no": "123"},
            ),
        ],
    )

    record = _build_null_record("AQ18", structured, doc_extraction, "sample.pdf")

    assert record.source_type == "null"
    assert record.null_category == "extraction_failure"
    assert record.blocks_approval is True


def test_null_record_marks_derived_dependency_missing():
    doc_extraction = DocumentExtraction(
        document_id="doc",
        filename="sample.pdf",
        sha256="sha",
        page_count=1,
        pages=[_page_meta(1)],
        tables=[_table_with_blank_static_cell()],
    )

    record = _build_null_record(
        "G5",
        _structured_page_1({"annual_volume_nos": None}),
        doc_extraction,
        "sample.pdf",
    )

    assert record.source_type == "null"
    assert record.null_category == "derived_dependency_missing"
    assert record.blocks_approval is False   # derived_dependency_missing is informational, not blocking


def test_cross_sheet_l_cells_only_grey_when_marked_unclear_logic():
    doc_extraction = DocumentExtraction(
        document_id="doc",
        filename="sample.pdf",
        sha256="sha",
        page_count=1,
        pages=[_page_meta(1)],
        tables=[_table_with_blank_static_cell()],
    )
    structured = MeridianExtractionResponse(
        document_id="doc",
        filename="sample.pdf",
        page_count=1,
        pages=[
            MeridianStructuredPage(
                page_number=1,
                page_type="gdc_estimation",
                capital_investments=[{"category": "Casting", "items": []}],
            )
        ],
    )

    assert CELL_SOURCE_TYPES.get("L17") != "unclear_logic"
    record = _build_null_record("L29", structured, doc_extraction, "sample.pdf")
    assert record.source_type == "null"
    assert record.null_category == "unclear_logic"


def test_formula_resolver_writes_formula_when_inputs_present():
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet["L23"] = 1
    sheet["M23"] = 47.25
    sheet["V5"] = 1.05
    extraction = SimpleNamespace(pages=[SimpleNamespace(capital_investments=[])])

    statuses = _resolve_formula_cells(sheet, extraction)

    assert statuses["N23"] == "formula"
    assert sheet["N23"].value == "=(M23*L23)*V5"
    assert sheet["N23"].fill.fgColor.rgb == FORMULA_FILL.fgColor.rgb


def test_n46_keeps_formula_when_m46_is_blank():
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet["V5"] = 1.05
    extraction = SimpleNamespace(
        pages=[
            SimpleNamespace(
                capital_investments=[
                    {
                        "category": "Others",
                        "items": [
                            {
                                "description": "Endoscope machine",
                                "total_cost_rs_lac": 8.0,
                            }
                        ],
                    }
                ],
            )
        ]
    )

    statuses = _resolve_formula_cells(sheet, extraction)

    assert statuses["N46"] == "formula"
    assert sheet["M46"].value is None
    assert sheet["N46"].value == "=M46*V5"


def test_formula_resolver_uses_pdf_fallback_when_inputs_missing():
    wb = openpyxl.Workbook()
    sheet = wb.active
    extraction = SimpleNamespace(
        pages=[
            SimpleNamespace(
                capital_investments=[
                    {
                        "category": "Sand core",
                        "items": [
                            {
                                "description": "Core shooting M/c",
                                "total_cost_rs_lac": 47.25,
                            }
                        ],
                    }
                ],
            )
        ]
    )

    statuses = _resolve_formula_cells(sheet, extraction)

    assert statuses["N23"] == "formula_fallback"
    assert sheet["N23"].value == 47.25


def test_formula_provenance_source_types_are_distinct():
    formula_record = _build_formula_record("N23", "=(M23*L23)*V5")
    fallback_record = _build_provenance_record(
        "N23",
        47.25,
        MeridianExtractionResponse(
            document_id="doc",
            filename="sample.pdf",
            page_count=1,
            pages=[MeridianStructuredPage(page_number=1, page_type="gdc_estimation")],
        ),
        DocumentExtraction(
            document_id="doc",
            filename="sample.pdf",
            sha256="sha",
            page_count=1,
            pages=[_page_meta(1)],
            tables=[],
        ),
        "sample.pdf",
        source_type_override="formula_fallback",
    )

    assert formula_record.source_type == "formula"
    assert formula_record.formula == "=(M23*L23)*V5"
    assert fallback_record.source_type == "formula_fallback"
    assert "Formula inputs were missing" in fallback_record.reason


def test_c6_formula_record_carries_formula_for_tracking_chips():
    record = _build_formula_record("C6", "=C5")

    assert record.excel_cell == "C6"
    assert record.source_type == "formula"
    assert record.formula == "=C5"


def test_i17_is_zeroed_when_formula_inputs_are_blank():
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet["I17"] = "=(D13)*I53*1.1"

    _apply_static_formula_cache_fixes(sheet)

    assert sheet["I17"].value == 0
