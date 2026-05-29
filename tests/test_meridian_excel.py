from pathlib import Path
import re
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from app.core.config import Settings, get_settings
from app.main import app, get_meridian_excel_fill_service, get_store
from app.models import DocumentExtraction
from app.services.meridian_excel import FIELD_TO_CELL, OUTPUT_DETAIL_ROWS, MeridianExcelFillService
from app.services.storage import DocumentStore


SAMPLE_EXTRACTION_PATH = Path("data/extractions/61f716326da04472.json")
SAMPLE_TEMPLATE_PATH = Path("samples/2425-328.xlsx")


def _load_sample_extraction() -> DocumentExtraction:
    return DocumentExtraction.model_validate_json(SAMPLE_EXTRACTION_PATH.read_text(encoding="utf-8"))


def _make_empty_input_template(tmp_path: Path) -> Path:
    workbook = load_workbook(SAMPLE_TEMPLATE_PATH)
    sheet = workbook["Input Sheet"]
    fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")

    target_cells = set(FIELD_TO_CELL.values())
    for row_number in OUTPUT_DETAIL_ROWS.values():
        target_cells.update({f"L{row_number}", f"M{row_number}", f"N{row_number}"})

    for cell_ref in target_cells:
        sheet[cell_ref] = None
        sheet[cell_ref].fill = fill

    path = tmp_path / "empty-input-template.xlsx"
    workbook.save(path)
    return path


def test_meridian_normalizer_extracts_core_fields() -> None:
    service = MeridianExcelFillService.__new__(MeridianExcelFillService)
    fields = service.normalize_meridian_fields(_load_sample_extraction())

    assert fields.fields["rfq_no"].value == "2425-328"
    assert fields.fields["customer"].value == "CTT"
    assert fields.fields["final_part_no"].value == "6511292"
    assert fields.fields["annual_volume"].value == 33000
    assert fields.fields["annual_volume_including_rejection"].value == 37950
    assert fields.fields["machined_part_weight_kg"].value == 2.96
    assert fields.fields["casting_weight_kg"].value == 3.8
    assert fields.fields["length_mm"].value == 242
    assert fields.fields["breadth_mm"].value == 211
    assert fields.fields["height_mm"].value == 134
    assert fields.fields["casting_process_category"].value == "GDC"

    sand_core_1 = fields.output_details["sand_core_1"]
    assert sand_core_1["cavities_loading"].value == 2
    assert sand_core_1["cycle_time_min"].value == 5
    assert sand_core_1["output_per_hr"].value == 24

    casting = fields.output_details["casting"]
    assert casting["cavities_loading"].value == 2
    assert casting["cycle_time_min"].value == 6
    assert casting["output_per_hr"].value == 20


def test_meridian_excel_service_fills_yellow_input_cells(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path, reference_pdf_path=SAMPLE_TEMPLATE_PATH)
    store = DocumentStore(settings)
    extraction = _load_sample_extraction()
    store.save_extraction(extraction)

    template_path = _make_empty_input_template(tmp_path)
    service = MeridianExcelFillService(settings, store)
    response = service.fill_input_sheet(extraction.document_id, template_path, template_path.name)

    output_path = store.excel_output_path(response.job_id)
    assert output_path.exists()

    with ZipFile(template_path) as template_zip, ZipFile(output_path) as output_zip:
        template_drawings = [name for name in template_zip.namelist() if name.startswith("xl/drawings/")]
        output_drawings = [name for name in output_zip.namelist() if name.startswith("xl/drawings/")]
        template_external_links = [name for name in template_zip.namelist() if name.startswith("xl/externalLinks/")]
        output_external_links = [name for name in output_zip.namelist() if name.startswith("xl/externalLinks/")]

        assert output_drawings == template_drawings
        assert output_external_links == template_external_links
        assert "xl/calcChain.xml" not in output_zip.namelist()

        sheet_path = service._sheet_xml_path(output_zip, "Input Sheet")
        sheet_xml = output_zip.read(sheet_path)
        root_start = sheet_xml[sheet_xml.find(b"<worksheet") : sheet_xml.find(b">", sheet_xml.find(b"<worksheet"))]
        ignorable = re.search(rb'Ignorable="([^"]+)"', root_start)
        if ignorable is not None:
            for prefix in ignorable.group(1).decode("utf-8").split():
                assert f"xmlns:{prefix}=".encode("utf-8") in root_start

    sheet = load_workbook(output_path, data_only=False)["Input Sheet"]
    assert sheet["C2"].value == "2425-328"
    assert sheet["C4"].value == "CTT"
    assert sheet["C5"].value == "6511292"
    assert sheet["G4"].value == 33000
    assert sheet["G5"].value == 37950
    assert sheet["G8"].value == 2.96
    assert sheet["G9"].value == 3.8
    assert sheet["G10"].value == 242
    assert sheet["G11"].value == 211
    assert sheet["G12"].value == 134
    assert sheet["K10"].value == "GDC"
    assert sheet["L12"].value == 2
    assert sheet["M12"].value == 5
    assert sheet["N12"].value == 24

    written_cells = {cell.cell for cell in response.report.written_cells if cell.status == "written"}
    assert {"C2", "C4", "C5", "G4", "G5", "L12", "M12", "N12"}.issubset(written_cells)
    assert any(field.field == "annual_volume_including_rejection" for field in response.report.validated_fields)


def test_meridian_excel_service_preserves_ignorable_namespace_prefixes(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path, reference_pdf_path=SAMPLE_TEMPLATE_PATH)
    store = DocumentStore(settings)
    extraction = _load_sample_extraction()
    store.save_extraction(extraction)

    service = MeridianExcelFillService(settings, store)
    response = service.fill_input_sheet(extraction.document_id, SAMPLE_TEMPLATE_PATH, SAMPLE_TEMPLATE_PATH.name)

    with ZipFile(store.excel_output_path(response.job_id)) as output_zip:
        sheet_path = service._sheet_xml_path(output_zip, "Input Sheet")
        sheet_xml = output_zip.read(sheet_path)
        root_start = sheet_xml[sheet_xml.find(b"<worksheet") : sheet_xml.find(b">", sheet_xml.find(b"<worksheet"))]
        ignorable = re.search(rb'Ignorable="([^"]+)"', root_start)
        assert ignorable is not None
        for prefix in ignorable.group(1).decode("utf-8").split():
            assert f"xmlns:{prefix}=".encode("utf-8") in root_start


def test_excel_fill_api_uploads_template_and_downloads_output(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path, reference_pdf_path=SAMPLE_TEMPLATE_PATH)
    store = DocumentStore(settings)
    extraction = _load_sample_extraction()
    store.save_extraction(extraction)

    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[get_store] = lambda: store
    app.dependency_overrides[get_meridian_excel_fill_service] = lambda: MeridianExcelFillService(settings, store)
    client = TestClient(app)

    try:
        template_path = _make_empty_input_template(tmp_path)
        with template_path.open("rb") as handle:
            response = client.post(
                f"/v1/documents/{extraction.document_id}/excel/input-sheet",
                files={
                    "template": (
                        template_path.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )

        assert response.status_code == 200
        payload = response.json()
        assert payload["document_id"] == extraction.document_id
        assert payload["download_url"] == f"/v1/excel/jobs/{payload['job_id']}/download"

        download_response = client.get(payload["download_url"])
        assert download_response.status_code == 200
        assert download_response.content.startswith(b"PK")
    finally:
        app.dependency_overrides.clear()
