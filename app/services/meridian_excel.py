from __future__ import annotations

import re
import posixpath
import tempfile
import uuid
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET

from fastapi import UploadFile
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet.worksheet import Worksheet

from app.core.config import Settings
from app.models import (
    DocumentExtraction,
    ExcelFillReport,
    ExcelFillResponse,
    ExcelValidatedField,
    ExcelWrittenCell,
    ExtractedTable,
    ExtractionWarning,
    FieldSource,
    NormalizedField,
    TableCell,
    WarningSeverity,
)
from app.services.errors import InvalidExcelTemplateError
from app.services.storage import DocumentStore


INPUT_SHEET_NAME = "Input Sheet"
YELLOW_RGB_VALUES = {"FFFFFF00", "FFFF00"}
SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
MARKUP_COMPATIBILITY_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
SPREADSHEET_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
SPREADSHEET_X14_NS = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
SPREADSHEET_X14AC_NS = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"
SPREADSHEET_XR_NS = "http://schemas.microsoft.com/office/spreadsheetml/2014/revision"
SPREADSHEET_XR2_NS = "http://schemas.microsoft.com/office/spreadsheetml/2015/revision2"
SPREADSHEET_XR3_NS = "http://schemas.microsoft.com/office/spreadsheetml/2016/revision3"
EXCEL_MAIN_NS = "http://schemas.microsoft.com/office/excel/2006/main"

ET.register_namespace("", SPREADSHEET_NS)
ET.register_namespace("r", OFFICE_REL_NS)
ET.register_namespace("mc", MARKUP_COMPATIBILITY_NS)
ET.register_namespace("xdr", SPREADSHEET_DRAWING_NS)
ET.register_namespace("x14", SPREADSHEET_X14_NS)
ET.register_namespace("x14ac", SPREADSHEET_X14AC_NS)
ET.register_namespace("xr", SPREADSHEET_XR_NS)
ET.register_namespace("xr2", SPREADSHEET_XR2_NS)
ET.register_namespace("xr3", SPREADSHEET_XR3_NS)
ET.register_namespace("xm", EXCEL_MAIN_NS)

FIELD_TO_CELL = {
    "rfq_no": "C2",
    "process": "C3",
    "process_scope": "D3",
    "customer": "C4",
    "final_part_no": "C5",
    "final_part_rev_no": "C6",
    "description": "C7",
    "alloy": "C8",
    "cavity": "C9",
    "geographical_segment": "C10",
    "incoterm": "D10",
    "annual_volume": "G4",
    "annual_volume_including_rejection": "G5",
    "machined_part_weight_kg": "G8",
    "casting_weight_kg": "G9",
    "length_mm": "G10",
    "breadth_mm": "G11",
    "height_mm": "G12",
    "casting_process_category": "K10",
}

OUTPUT_DETAIL_ROWS = {
    "sand_core_1": 12,
    "sand_core_2": 13,
    "sand_core_3": 14,
    "sand_core_4": 15,
    "decore": 16,
    "casting": 17,
    "cutting_1": 18,
    "cutting_2": 19,
    "linishing": 20,
    "shot_blasting": 21,
}

OUTPUT_OPERATION_ALIASES = {
    "sand core 1": "sand_core_1",
    "sand core 2": "sand_core_2",
    "sand core 3": "sand_core_3",
    "sand core 4": "sand_core_4",
    "decore": "decore",
    "casting": "casting",
    "cutting 1": "cutting_1",
    "cutting 2": "cutting_2",
    "linishing": "linishing",
    "shot blasting": "shot_blasting",
}

LABEL_ALIASES = {
    "rfq no": "rfq_no",
    "customer": "customer",
    "final part no": "final_part_no",
    "final part rev no": "final_part_rev_no",
    "description": "description",
    "alloy": "alloy",
    "annual volume": "annual_volume",
    "annual volume nos": "annual_volume",
    "annual volume including rejection": "annual_volume_including_rejection",
    "annual volume including rejection nos": "annual_volume_including_rejection",
    "annual volume incl rejection": "annual_volume_including_rejection",
    "annual volume incl rejection nos": "annual_volume_including_rejection",
    "machined part wt": "machined_part_weight_kg",
    "machined part wt kg": "machined_part_weight_kg",
    "casting weight": "casting_weight_kg",
    "casting weight kg": "casting_weight_kg",
    "lbh": "lbh_mm",
    "lbh mm": "lbh_mm",
    "casting category": "casting_category",
    "no of casting cells planned": "cavity",
}

ANCHOR_LABELS = {
    "B2": "RFQ No.",
    "B4": "Customer",
    "B5": "Final Part No.",
    "B7": "Description",
    "B8": "Alloy",
    "F8": "Machined part wt (kg)",
    "F9": "Casting Weight (Kg)",
}


@dataclass
class MeridianFields:
    fields: dict[str, NormalizedField]
    output_details: dict[str, dict[str, NormalizedField]]
    warnings: list[ExtractionWarning]


class MeridianExcelFillService:
    def __init__(self, settings: Settings, store: DocumentStore) -> None:
        self.settings = settings
        self.store = store

    async def save_template_to_temp(self, upload: UploadFile) -> Path:
        filename = upload.filename or "template.xlsx"
        if Path(filename).suffix.lower() != ".xlsx":
            raise InvalidExcelTemplateError("Only .xlsx templates are supported.", details={"filename": filename})

        written = 0
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        tmp_path = Path(tmp.name)
        try:
            while chunk := await upload.read(1024 * 1024):
                written += len(chunk)
                if written > self.settings.max_upload_bytes:
                    raise InvalidExcelTemplateError(
                        "Uploaded Excel template exceeds the configured size limit.",
                        details={"max_upload_bytes": self.settings.max_upload_bytes},
                    )
                tmp.write(chunk)
        finally:
            tmp.close()

        if written == 0:
            tmp_path.unlink(missing_ok=True)
            raise InvalidExcelTemplateError("Uploaded Excel template is empty.")
        return tmp_path

    def fill_input_sheet(
        self,
        document_id: str,
        template_path: Path,
        template_filename: str,
    ) -> ExcelFillResponse:
        extraction = self.store.load_extraction(document_id)
        meridian_fields = self.normalize_meridian_fields(extraction)
        job_id = uuid.uuid4().hex[:16]

        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", category=UserWarning, module=r"openpyxl\..*")
            workbook = load_workbook(template_path)
        if INPUT_SHEET_NAME not in workbook.sheetnames:
            raise InvalidExcelTemplateError(
                "Excel template does not contain the required Input Sheet.",
                details={"required_sheet": INPUT_SHEET_NAME},
            )
        sheet = workbook[INPUT_SHEET_NAME]

        report = ExcelFillReport(
            job_id=job_id,
            document_id=document_id,
            template_filename=template_filename,
            output_filename="filled.xlsx",
            warnings=list(meridian_fields.warnings),
        )

        self._validate_anchor_labels(sheet, report)
        yellow_cells = self._yellow_cells(sheet)
        cell_updates: dict[str, Any] = {}
        self._write_field_cells(meridian_fields, yellow_cells, report, cell_updates)
        self._write_output_detail_cells(meridian_fields, yellow_cells, report, cell_updates)

        output_tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        output_path = Path(output_tmp.name)
        output_tmp.close()
        self._write_xlsx_with_cell_updates(template_path, INPUT_SHEET_NAME, cell_updates, output_path)
        self.store.save_excel_output(job_id, output_path, report)
        output_path.unlink(missing_ok=True)

        return ExcelFillResponse(
            job_id=job_id,
            document_id=document_id,
            output_filename=report.output_filename,
            download_url=f"/v1/excel/jobs/{job_id}/download",
            report=report,
        )

    def normalize_meridian_fields(self, extraction: DocumentExtraction) -> MeridianFields:
        fields: dict[str, NormalizedField] = {}
        output_details: dict[str, dict[str, NormalizedField]] = {}
        warnings: list[ExtractionWarning] = []

        for table in sorted(extraction.tables, key=lambda item: self._table_priority(item)):
            rows = list(table.rows)
            for index, row in enumerate(rows):
                cells = [cell for cell in row.cells if cell.text]
                row_items = [(cell, self._clean_text(cell.text)) for cell in cells]
                self._extract_label_pairs(table, row_items, fields)
                next_row_items: list[tuple[TableCell, str]] | None = None
                if index + 1 < len(rows):
                    next_row_items = [
                        (cell, self._clean_text(cell.text))
                        for cell in rows[index + 1].cells
                        if cell.text
                    ]
                self._extract_output_detail_row(table, row_items, output_details, next_row_items)

            if table.title and "gravity die casting" in table.title.lower():
                self._set_field(fields, "process", "GDC with core", None, confidence=0.75)
                self._set_field(fields, "casting_process_category", "GDC", None, confidence=0.75)
            if table.title and "machining" in table.title.lower():
                self._set_field(fields, "process_scope", "With Machining", None, confidence=0.75)

        self._split_lbh(fields)
        self._calculate_derived_fields(fields, output_details)
        self._warn_missing_core_fields(fields, warnings)
        return MeridianFields(fields=fields, output_details=output_details, warnings=warnings)

    def _extract_label_pairs(
        self,
        table: ExtractedTable,
        row_items: list[tuple[TableCell, str]],
        fields: dict[str, NormalizedField],
    ) -> None:
        for index, (cell, text) in enumerate(row_items):
            label = self._normalize_label(text)
            field_name = LABEL_ALIASES.get(label)
            if not field_name or index + 1 >= len(row_items):
                continue
            value_cell, value = row_items[index + 1]
            if value in LABEL_ALIASES or not value:
                continue
            source = FieldSource(
                page_number=table.page_number,
                table_id=table.table_id,
                row=cell.row,
                column=value_cell.column,
                text=value_cell.text,
            )
            parsed_value = self._parse_field_value(field_name, value)
            self._set_field(fields, field_name, parsed_value, source)

    def _extract_output_detail_row(
        self,
        table: ExtractedTable,
        row_items: list[tuple[TableCell, str]],
        output_details: dict[str, dict[str, NormalizedField]],
        next_row_items: list[tuple[TableCell, str]] | None = None,
    ) -> None:
        if not row_items:
            return

        operation_cell, operation_text = row_items[0]
        if "output/ m/c details" in operation_text.lower():
            operation_text = "Sand core 1"
            if next_row_items:
                row_items = [(operation_cell, operation_text), *next_row_items]
        operation_key = OUTPUT_OPERATION_ALIASES.get(self._normalize_operation(operation_text))
        if not operation_key:
            return

        values_by_column = {cell.column: (cell, text) for cell, text in row_items}
        detail = output_details.setdefault(operation_key, {})
        for field_name, column in (
            ("cavities_loading", 3),
            ("cycle_time_min", 4),
            ("output_per_hr", 5),
        ):
            if column not in values_by_column:
                continue
            value_cell, text = values_by_column[column]
            source = FieldSource(
                page_number=table.page_number,
                table_id=table.table_id,
                row=operation_cell.row,
                column=value_cell.column,
                text=value_cell.text,
            )
            detail.setdefault(field_name, NormalizedField(value=self._parse_number(text), source=source, confidence=0.9))

    def _split_lbh(self, fields: dict[str, NormalizedField]) -> None:
        lbh = fields.get("lbh_mm")
        if not lbh or not isinstance(lbh.value, str):
            return
        numbers = [self._parse_number(match) for match in re.findall(r"\d+(?:\.\d+)?", lbh.value)]
        if len(numbers) < 3:
            return
        for name, value in zip(("length_mm", "breadth_mm", "height_mm"), numbers[:3], strict=True):
            self._set_field(fields, name, value, lbh.source, confidence=0.92)

    def _calculate_derived_fields(
        self,
        fields: dict[str, NormalizedField],
        output_details: dict[str, dict[str, NormalizedField]],
    ) -> None:
        annual_volume = fields.get("annual_volume")
        if annual_volume and "annual_volume_including_rejection" not in fields:
            calculated = round(float(annual_volume.value) * 1.15)
            self._set_field(fields, "annual_volume_including_rejection", calculated, annual_volume.source, confidence=0.75)

        for detail in output_details.values():
            if "output_per_hr" in detail:
                continue
            cavities = detail.get("cavities_loading")
            cycle_time = detail.get("cycle_time_min")
            if cavities and cycle_time and float(cycle_time.value) != 0:
                detail["output_per_hr"] = NormalizedField(
                    value=round((60 / float(cycle_time.value)) * float(cavities.value), 6),
                    source=cycle_time.source,
                    confidence=0.75,
                )

    def _write_field_cells(
        self,
        meridian_fields: MeridianFields,
        yellow_cells: set[str],
        report: ExcelFillReport,
        cell_updates: dict[str, Any],
    ) -> None:
        for field_name, cell_ref in FIELD_TO_CELL.items():
            field = meridian_fields.fields.get(field_name)
            if field is None:
                report.warnings.append(self._warning("missing_extracted_field", f"Missing extracted field '{field_name}'.", field=field_name, cell=cell_ref))
                continue
            if not self._is_yellow_target(cell_ref, yellow_cells):
                report.warnings.append(self._warning("target_cell_not_yellow", f"Target cell {cell_ref} is not yellow; skipped.", field=field_name, cell=cell_ref))
                report.written_cells.append(ExcelWrittenCell(field=field_name, cell=cell_ref, value=field.value, status="skipped", source=field.source))
                continue

            cell_updates[cell_ref] = field.value
            report.written_cells.append(ExcelWrittenCell(field=field_name, cell=cell_ref, value=field.value, status="written", source=field.source))

            if field_name == "annual_volume_including_rejection":
                self._validate_annual_volume(meridian_fields, field.value, cell_ref, report)

    def _write_output_detail_cells(
        self,
        meridian_fields: MeridianFields,
        yellow_cells: set[str],
        report: ExcelFillReport,
        cell_updates: dict[str, Any],
    ) -> None:
        columns = {
            "cavities_loading": "L",
            "cycle_time_min": "M",
            "output_per_hr": "N",
        }
        for operation_key, row_number in OUTPUT_DETAIL_ROWS.items():
            detail = meridian_fields.output_details.get(operation_key, {})
            for field_name, column in columns.items():
                cell_ref = f"{column}{row_number}"
                field = detail.get(field_name)
                report_field = f"output_details.{operation_key}.{field_name}"
                if field is None:
                    report.warnings.append(self._warning("missing_extracted_field", f"Missing extracted field '{report_field}'.", field=report_field, cell=cell_ref))
                    continue
                if not self._is_yellow_target(cell_ref, yellow_cells):
                    report.warnings.append(self._warning("target_cell_not_yellow", f"Target cell {cell_ref} is not yellow; skipped.", field=report_field, cell=cell_ref))
                    continue
                cell_updates[cell_ref] = field.value
                report.written_cells.append(ExcelWrittenCell(field=report_field, cell=cell_ref, value=field.value, status="written", source=field.source))
            self._validate_output_detail(operation_key, row_number, detail, report)

    def _write_xlsx_with_cell_updates(
        self,
        template_path: Path,
        sheet_name: str,
        cell_updates: dict[str, Any],
        output_path: Path,
    ) -> None:
        with zipfile.ZipFile(template_path, "r") as source:
            sheet_path = self._sheet_xml_path(source, sheet_name)
            patched_sheet = self._patch_sheet_xml(source.read(sheet_path), cell_updates)
            patched_rels = self._remove_calc_chain_relationship(source.read("xl/_rels/workbook.xml.rels"))
            patched_content_types = self._remove_calc_chain_content_type(source.read("[Content_Types].xml"))

            with zipfile.ZipFile(output_path, "w") as destination:
                for item in source.infolist():
                    if item.filename == "xl/calcChain.xml":
                        continue
                    if item.filename == sheet_path:
                        destination.writestr(item, patched_sheet)
                    elif item.filename == "xl/_rels/workbook.xml.rels":
                        destination.writestr(item, patched_rels)
                    elif item.filename == "[Content_Types].xml":
                        destination.writestr(item, patched_content_types)
                    else:
                        destination.writestr(item, source.read(item.filename))

    def _sheet_xml_path(self, workbook: zipfile.ZipFile, sheet_name: str) -> str:
        workbook_root = ET.fromstring(workbook.read("xl/workbook.xml"))
        rels_root = ET.fromstring(workbook.read("xl/_rels/workbook.xml.rels"))
        sheet_rel_id: str | None = None
        for sheet in workbook_root.findall(f".//{{{SPREADSHEET_NS}}}sheet"):
            if sheet.attrib.get("name") == sheet_name:
                sheet_rel_id = sheet.attrib.get(f"{{{OFFICE_REL_NS}}}id")
                break
        if sheet_rel_id is None:
            raise InvalidExcelTemplateError(
                "Excel template does not contain the required Input Sheet.",
                details={"required_sheet": sheet_name},
            )

        for relationship in rels_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship"):
            if relationship.attrib.get("Id") != sheet_rel_id:
                continue
            target = relationship.attrib["Target"]
            if target.startswith("/"):
                return target.lstrip("/")
            return posixpath.normpath(posixpath.join("xl", target))

        raise InvalidExcelTemplateError(
            "Input Sheet relationship could not be found in the Excel template.",
            details={"required_sheet": sheet_name},
        )

    def _patch_sheet_xml(self, sheet_xml: bytes, cell_updates: dict[str, Any]) -> bytes:
        root = ET.fromstring(sheet_xml)
        sheet_data = root.find(f"{{{SPREADSHEET_NS}}}sheetData")
        if sheet_data is None:
            raise InvalidExcelTemplateError("Input Sheet XML does not contain sheetData.")

        rows_by_number = {
            int(row.attrib["r"]): row
            for row in sheet_data.findall(f"{{{SPREADSHEET_NS}}}row")
            if row.attrib.get("r", "").isdigit()
        }

        for cell_ref, value in cell_updates.items():
            row_number, column_number = coordinate_to_tuple(cell_ref)
            row = rows_by_number.get(row_number)
            if row is None:
                row = ET.Element(f"{{{SPREADSHEET_NS}}}row", {"r": str(row_number)})
                self._insert_row(sheet_data, row)
                rows_by_number[row_number] = row

            cell = self._find_or_create_cell(row, cell_ref, column_number)
            self._set_cell_value(cell, value)

        patched_xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        return self._restore_ignorable_namespace_declarations(sheet_xml, patched_xml)

    def _restore_ignorable_namespace_declarations(self, original_xml: bytes, patched_xml: bytes) -> bytes:
        original_root = self._root_start_tag(original_xml)
        ignorable_match = re.search(rb"(?:\w+:)?Ignorable=\"([^\"]+)\"", original_root)
        if not ignorable_match:
            return patched_xml

        original_namespaces = {
            match.group(1).decode("utf-8"): match.group(2)
            for match in re.finditer(rb"xmlns:([A-Za-z0-9_]+)=\"([^\"]+)\"", original_root)
        }
        patched_root = self._root_start_tag(patched_xml)
        declarations_to_add = []
        for prefix in ignorable_match.group(1).decode("utf-8").split():
            if f"xmlns:{prefix}=".encode("utf-8") in patched_root:
                continue
            namespace_uri = original_namespaces.get(prefix)
            if namespace_uri:
                declarations_to_add.append(b' xmlns:' + prefix.encode("utf-8") + b'="' + namespace_uri + b'"')

        if not declarations_to_add:
            return patched_xml

        insert_at = patched_xml.find(b">", patched_xml.find(b"<worksheet"))
        if insert_at == -1:
            return patched_xml
        return patched_xml[:insert_at] + b"".join(declarations_to_add) + patched_xml[insert_at:]

    def _root_start_tag(self, xml: bytes) -> bytes:
        start = xml.find(b"<worksheet")
        end = xml.find(b">", start)
        if start == -1 or end == -1:
            return b""
        return xml[start : end + 1]

    def _insert_row(self, sheet_data: ET.Element, row: ET.Element) -> None:
        row_number = int(row.attrib["r"])
        for index, existing_row in enumerate(list(sheet_data)):
            existing_number = int(existing_row.attrib.get("r", "0") or 0)
            if existing_number > row_number:
                sheet_data.insert(index, row)
                return
        sheet_data.append(row)

    def _find_or_create_cell(self, row: ET.Element, cell_ref: str, column_number: int) -> ET.Element:
        for cell in row.findall(f"{{{SPREADSHEET_NS}}}c"):
            if cell.attrib.get("r") == cell_ref:
                return cell

        cell = ET.Element(f"{{{SPREADSHEET_NS}}}c", {"r": cell_ref})
        for index, existing_cell in enumerate(row.findall(f"{{{SPREADSHEET_NS}}}c")):
            existing_ref = existing_cell.attrib.get("r", "")
            existing_column = coordinate_to_tuple(existing_ref)[1] if existing_ref else 0
            if existing_column > column_number:
                row.insert(index, cell)
                return cell
        row.append(cell)
        return cell

    def _set_cell_value(self, cell: ET.Element, value: Any) -> None:
        cell_ref = cell.attrib.get("r")
        if not cell_ref:
            raise InvalidExcelTemplateError("A target Excel cell is missing its cell reference.")
        style_id = cell.attrib.get("s")
        cell.clear()
        cell.attrib["r"] = cell_ref
        if style_id is not None:
            cell.attrib["s"] = style_id

        if isinstance(value, bool):
            cell.attrib["t"] = "b"
            value_element = ET.SubElement(cell, f"{{{SPREADSHEET_NS}}}v")
            value_element.text = "1" if value else "0"
            return

        if isinstance(value, (int, float)):
            value_element = ET.SubElement(cell, f"{{{SPREADSHEET_NS}}}v")
            value_element.text = str(value)
            return

        cell.attrib["t"] = "inlineStr"
        inline_string = ET.SubElement(cell, f"{{{SPREADSHEET_NS}}}is")
        text = ET.SubElement(inline_string, f"{{{SPREADSHEET_NS}}}t")
        text.text = str(value)

    def _remove_calc_chain_relationship(self, rels_xml: bytes) -> bytes:
        root = ET.fromstring(rels_xml)
        for relationship in list(root):
            if relationship.attrib.get("Type", "").endswith("/calcChain"):
                root.remove(relationship)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    def _remove_calc_chain_content_type(self, content_types_xml: bytes) -> bytes:
        root = ET.fromstring(content_types_xml)
        for override in list(root):
            if override.attrib.get("PartName") == "/xl/calcChain.xml":
                root.remove(override)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    def _validate_annual_volume(self, meridian_fields: MeridianFields, extracted_value: Any, cell_ref: str, report: ExcelFillReport) -> None:
        annual_volume = meridian_fields.fields.get("annual_volume")
        if annual_volume is None:
            report.validated_fields.append(
                ExcelValidatedField(field="annual_volume_including_rejection", cell=cell_ref, calculated_value=None, extracted_value=extracted_value, status="missing_inputs")
            )
            return
        calculated = round(float(annual_volume.value) * 1.15)
        status = "validated" if self._numbers_close(calculated, extracted_value) else "mismatch"
        report.validated_fields.append(
            ExcelValidatedField(
                field="annual_volume_including_rejection",
                cell=cell_ref,
                calculated_value=calculated,
                extracted_value=extracted_value,
                status=status,
            )
        )

    def _validate_output_detail(
        self,
        operation_key: str,
        row_number: int,
        detail: dict[str, NormalizedField],
        report: ExcelFillReport,
    ) -> None:
        cavities = detail.get("cavities_loading")
        cycle_time = detail.get("cycle_time_min")
        output = detail.get("output_per_hr")
        if not cavities or not cycle_time or not output:
            return
        if float(cycle_time.value) == 0:
            return
        calculated = round((60 / float(cycle_time.value)) * float(cavities.value), 6)
        status = "validated" if self._numbers_close(calculated, output.value) else "mismatch"
        report.validated_fields.append(
            ExcelValidatedField(
                field=f"output_details.{operation_key}.output_per_hr",
                cell=f"N{row_number}",
                calculated_value=calculated,
                extracted_value=output.value,
                status=status,
            )
        )

    def _validate_anchor_labels(self, sheet: Worksheet, report: ExcelFillReport) -> None:
        for cell_ref, expected in ANCHOR_LABELS.items():
            actual = self._clean_text(str(sheet[cell_ref].value or ""))
            if self._normalize_label(actual) != self._normalize_label(expected):
                report.warnings.append(
                    self._warning(
                        "template_anchor_mismatch",
                        f"Expected '{expected}' at {cell_ref}, found '{actual}'.",
                        cell=cell_ref,
                        expected=expected,
                        actual=actual,
                    )
                )

    def _yellow_cells(self, sheet: Worksheet) -> set[str]:
        yellow: set[str] = set()
        for row in sheet.iter_rows():
            for cell in row:
                if self._is_yellow_cell(cell):
                    yellow.add(cell.coordinate)
        return yellow

    def _is_yellow_cell(self, cell: Cell) -> bool:
        fill = cell.fill
        if fill is None or fill.patternType != "solid":
            return False
        color = fill.fgColor
        rgb = color.rgb
        if rgb is not None:
            return str(rgb).upper() in YELLOW_RGB_VALUES
        return False

    def _is_yellow_target(self, cell_ref: str, yellow_cells: set[str]) -> bool:
        return cell_ref in yellow_cells

    def _table_priority(self, table: ExtractedTable) -> tuple[int, int]:
        page_priority = {1: 0, 6: 1, 7: 2, 3: 3, 4: 4, 5: 5}
        return (page_priority.get(table.page_number, 99), table.page_number)

    def _set_field(
        self,
        fields: dict[str, NormalizedField],
        name: str,
        value: Any,
        source: FieldSource | None,
        *,
        confidence: float = 0.9,
    ) -> None:
        if value in (None, "") or name in fields:
            return
        fields[name] = NormalizedField(value=value, source=source, confidence=confidence)

    def _warn_missing_core_fields(self, fields: dict[str, NormalizedField], warnings: list[ExtractionWarning]) -> None:
        for field_name in FIELD_TO_CELL:
            if field_name not in fields:
                warnings.append(self._warning("missing_extracted_field", f"Missing extracted field '{field_name}'.", field=field_name))

    def _clean_text(self, value: str) -> str:
        return re.sub(r"\s+", " ", value).strip()

    def _normalize_label(self, value: str) -> str:
        cleaned = value.lower().replace("&", "and")
        cleaned = re.sub(r"\([^)]*\)", "", cleaned)
        cleaned = re.sub(r"[^a-z0-9]+", " ", cleaned)
        return re.sub(r"\s+", " ", cleaned).strip()

    def _normalize_operation(self, value: str) -> str:
        cleaned = self._clean_text(value).lower()
        cleaned = cleaned.replace("output/ m/c details:", "")
        cleaned = re.sub(r"[^a-z0-9]+", " ", cleaned)
        return re.sub(r"\s+", " ", cleaned).strip()

    def _parse_field_value(self, field_name: str, value: str) -> Any:
        if field_name in {
            "annual_volume",
            "annual_volume_including_rejection",
            "machined_part_weight_kg",
            "casting_weight_kg",
            "cavity",
        }:
            return self._parse_number(value)
        return value

    def _parse_number(self, value: Any) -> int | float | str:
        if isinstance(value, (int, float)):
            return value
        cleaned = str(value).replace(",", "").strip()
        match = re.search(r"-?\d+(?:\.\d+)?", cleaned)
        if not match:
            return value
        number = float(match.group(0))
        return int(number) if number.is_integer() else number

    def _numbers_close(self, left: Any, right: Any, *, tolerance: float = 0.01) -> bool:
        try:
            return abs(float(left) - float(right)) <= tolerance
        except (TypeError, ValueError):
            return False

    def _warning(self, code: str, message: str, **details: Any) -> ExtractionWarning:
        return ExtractionWarning(code=code, message=message, severity=WarningSeverity.warning, details=details)
